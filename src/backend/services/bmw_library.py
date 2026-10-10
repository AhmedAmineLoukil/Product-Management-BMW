"""Local-only adapters for the supplied BMW files.

The files stay in ignored data/raw/. This module returns one vehicle/market
scope at a time and never writes a derived customer-data file.
"""
from __future__ import annotations

import csv
import hashlib
import math
import re
from collections import OrderedDict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
RAW = ROOT / 'data' / 'raw'
FILES = {
    'feedback': 'F70_feedback_hackathon.xlsx',
    'studies': 'F70_G60_G68_G70_customer_studies.xlsx',
    'sales': 'sales_volumes.xlsx',
    'equipment': 'f70_equipment_curated.csv',
    'catalog': 'F70_OptionList.PDF',
}
VEHICLES = {
    'F70': 'BMW 1 Series F70',
    'G60': 'BMW 5 Series G60',
    'G68': 'BMW 5 Series G68',
    'G70': 'BMW 7 Series G70',
}
MARKETS = {
    'GB': 'United Kingdom', 'JP': 'Japan', 'SE': 'Sweden',
    'FR': 'France', 'AU': 'Australia', 'DK': 'Denmark',
    'FI': 'Finland', 'KR': 'South Korea', 'ZA': 'South Africa',
    'DE': 'Germany', 'EU': 'Europe', 'CN': 'China',
    'US': 'United States', 'RoW': 'Rest of World',
}
DATASET_NAMES = {
    'feedback': 'Customer verbatims', 'studies': 'Customer study',
    'sales': 'Sales volumes', 'equipment': 'Equipment list',
}


def _id(prefix: str, *parts: object) -> str:
    digest = hashlib.sha256('\n'.join(map(str, parts)).encode()).hexdigest()[:16]
    return f'BMW-{prefix}-{digest}'


def _code(value: object) -> str:
    match = re.search(r'\b(F70|G60|G68|G70)\b', str(value or '').upper())
    return match.group(1) if match else ''


def _number(value: object) -> float | None:
    return float(value) if isinstance(value, (int, float)) and math.isfinite(value) else None


def _record(kind: str, model: str, market: str, source: str, sheet: str,
            row: int, content: str, fields: dict, *, identity: object,
            category: str = '') -> dict:
    return {
        'id': _id(kind.upper(), model, market, source, sheet, identity),
        'content': content, 'row': row, 'sheet': sheet, 'category': category,
        'fields': {
            'source_file': source, 'source_sheet': sheet, 'source_row': row,
            'source_vehicle_code': model, 'source_market_code': market,
            'source_market_label': MARKETS[market], 'source_type': kind,
            **fields,
        },
    }


def _feedback() -> dict[tuple[str, str], list[dict]]:
    path = RAW / FILES['feedback']
    if not path.exists():
        return {}
    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book.active
    rows = sheet.values
    header = next(rows)
    grouped: OrderedDict[str, dict] = OrderedDict()
    for row_number, values in enumerate(rows, start=2):
        item = dict(zip(header, values))
        original_id = str(item.get('ID') or '').strip()
        text = str(item.get('Customer Feedback') or '').strip()
        market = str(item.get('Country') or '').strip().upper()
        model = _code(item.get('Derivate (e-code)'))
        if not original_id or not text or model != 'F70' or market not in MARKETS:
            raise ValueError(f'Feedback row {row_number} has missing or unsupported identity, text, vehicle, or market.')
        annotation = {
            'row': row_number, 'feedback_type': item.get('Feedback Type'),
            'category_level_2': item.get('Vfc level2 Name'),
            'category_level_2_3': item.get('Vfc level2 3 Name'),
            'category_level_3': item.get('Vfc level3 Name'),
        }
        if original_id not in grouped:
            grouped[original_id] = {
                'text': text, 'market': market, 'model': model,
                'first_row': row_number, 'source': item.get('Source'),
                'brand': item.get('Brand'), 'sales_description': item.get('Sales Description'),
                'engine_type': item.get('Engine Type'), 'annotations': [],
            }
        group = grouped[original_id]
        if (group['text'], group['market'], group['model']) != (text, market, model):
            raise ValueError(f'Feedback ID {original_id} has conflicting text, vehicle, or market across rows.')
        group['annotations'].append(annotation)
    book.close()
    result: dict[tuple[str, str], list[dict]] = {}
    for original_id, group in grouped.items():
        model, market = group['model'], group['market']
        record = _record('feedback', model, market, path.name, sheet.title,
                         group['first_row'], group['text'], {
                             '_originalId': original_id, 'original_row_count': len(group['annotations']),
                             'annotations': group['annotations'], 'source': group['source'],
                             'brand': group['brand'], 'sales_description': group['sales_description'],
                             'engine_type': group['engine_type'],
                         }, identity=original_id,
                         category=str(group['annotations'][0].get('category_level_2') or ''))
        result.setdefault((model, market), []).append(record)
    return result


def _studies() -> dict[tuple[str, str], list[dict]]:
    path = RAW / FILES['studies']
    if not path.exists():
        return {}
    book = load_workbook(path, read_only=True, data_only=True)
    result: dict[tuple[str, str], list[dict]] = {}
    if 'CN_EU_2025' in book:
        sheet = book['CN_EU_2025']
        rows = list(sheet.values)
        models = rows[2]
        market = ''
        scopes: dict[int, tuple[str, str]] = {}
        for column in range(1, len(models)):
            if rows[1][column]:
                market = {'EU': 'EU', 'China': 'CN'}.get(str(rows[1][column]), '')
            model = _code(models[column])
            if model and market:
                scopes[column] = model, market
        category = str(models[0] or '')
        for row_number, row in enumerate(rows[3:], start=4):
            if str(row[0] or '').strip().lower() != 'mean':
                if row[0]:
                    category = str(row[0]).strip()
                continue
            for column, (model, market) in scopes.items():
                value = _number(row[column] if column < len(row) else None)
                if value is None:
                    continue
                fields = {
                    'category': category, 'measure': 'mean rating', 'value': value,
                    'scale': 'Not stated in workbook', 'source_column': column + 1,
                    'source_model_label': str(models[column]), 'study_year': 2025,
                }
                result.setdefault((model, market), []).append(_record(
                    'study', model, market, path.name, sheet.title, row_number,
                    f'{category}: mean rating {value:g} (source scale not stated)',
                    fields, identity=f'R{row_number}C{column+1}', category=category))
    if 'US_2025' in book:
        sheet = book['US_2025']
        rows = list(sheet.values)
        models = rows[0]
        scopes = {column: _code(models[column]) for column in range(2, len(models)) if _code(models[column])}
        category = str(models[0] or '')
        for row_number, row in enumerate(rows[1:], start=2):
            has_value = any(_number(row[col] if col < len(row) else None) is not None for col in scopes)
            if not has_value:
                if row[0]:
                    category = str(row[0]).strip()
                continue
            response = str(row[0] or '').strip()
            measure = str(row[1] or '').strip()
            for column, model in scopes.items():
                value = _number(row[column] if column < len(row) else None)
                if value is None:
                    continue
                fields = {
                    'category': category, 'response': response, 'measure': measure,
                    'value': value, 'unit': 'fraction of column; not a mean rating',
                    'source_column': column + 1, 'source_model_label': str(models[column]),
                    'study_year': 2025,
                }
                result.setdefault((model, 'US'), []).append(_record(
                    'study', model, 'US', path.name, sheet.title, row_number,
                    f'{category} — {response}: {value:g} of column ({measure})',
                    fields, identity=f'R{row_number}C{column+1}', category=category))
    book.close()
    return result


def _sales() -> dict[tuple[str, str], list[dict]]:
    path = RAW / FILES['sales']
    if not path.exists():
        return {}
    book = load_workbook(path, read_only=True, data_only=True)
    result: dict[tuple[str, str], list[dict]] = {}
    for sheet in book:
        source_scope = sheet.title
        targets = ['G60', 'G68'] if source_scope == 'G60_G68' else [source_scope]
        if any(model not in VEHICLES for model in targets):
            continue
        rows = sheet.values
        header = next(rows)
        for row_number, values in enumerate(rows, start=2):
            item = dict(zip(header, values))
            market = str(item.get('market_code') or '').strip()
            if market not in MARKETS:
                raise ValueError(f'Sales row {row_number} in {source_scope} has an unsupported market.')
            for year in (2024, 2025, 2030):
                value = _number(item.get(f'volume_{year}'))
                if value is None:
                    continue
                for model in targets:
                    shared = source_scope == 'G60_G68'
                    note = ('Combined G60/G68 series total; it cannot be attributed to either model alone. '
                            if shared else '')
                    if year == 2030:
                        note += '2030 is future-looking in the supplied file; its forecast basis is not stated.'
                    fields = {
                        'market_original': item.get('market'), 'market_code_original': market,
                        'year': year, 'volume': value, 'source_model_scope': source_scope,
                        'combined_model_scope': shared, 'future_value': year == 2030,
                        'interpretation_note': note,
                    }
                    record = _record('sales', model, market, path.name, sheet.title,
                                     row_number, f'{MARKETS[market]} · {source_scope} · {year}: {value:g} units. {note}'.strip(),
                                     fields, identity=f'{market}-{year}')
                    result.setdefault((model, market), []).append(record)
    book.close()
    return result


def _equipment() -> dict[tuple[str, str], list[dict]]:
    path, pdf = RAW / FILES['equipment'], RAW / FILES['catalog']
    if not path.exists() or not pdf.exists():
        return {}
    import pymupdf

    document = pymupdf.open(pdf)
    pages = {i + 1: page.get_text() for i, page in enumerate(document)}
    document.close()
    result: list[dict] = []
    for page, text in pages.items():
        result.append(_record('catalog', 'F70', 'DE', pdf.name, 'PDF page', page,
                              text, {'catalog_page': page, 'availability': 'See original table; not auto-classified'},
                              identity=page, category='Catalog page'))
    with path.open(encoding='utf-8-sig', newline='') as stream:
        for index, item in enumerate(csv.DictReader(stream), start=2):
            page = int(item['page'])
            quote = item['source_text_match'].strip()
            if page not in pages or not quote or quote not in pages[page]:
                raise ValueError(f'Equipment CSV row {index} does not match catalog PDF page {page}.')
            name = item['feature_name_en'].strip() or item['feature_name_original'].strip()
            fields = {**item, 'source_pdf': pdf.name, 'source_pdf_page': page,
                      'original_csv_row': index,
                      'review_note': 'Curated classification is an initial manual review; check catalog symbols and variant restrictions.'}
            result.append(_record('equipment', 'F70', 'DE', path.name, 'Curated CSV', index,
                                  f'{name}: {item["availability"]}. PDF page {page}: {quote}',
                                  fields, identity=index, category=item['category']))
    return {('F70', 'DE'): result}


def _all() -> dict[str, dict[tuple[str, str], list[dict]]]:
    return {'feedback': _feedback(), 'studies': _studies(), 'sales': _sales(), 'equipment': _equipment()}


def manifest() -> dict:
    available = {key: (RAW / name).exists() for key, name in FILES.items()}
    data = _all()
    scopes: dict[tuple[str, str], dict] = {}
    for kind, mapping in data.items():
        for (model, market), records in mapping.items():
            entry = scopes.setdefault((model, market), {
                'model_code': model, 'vehicle': VEHICLES[model],
                'market_code': market, 'market': MARKETS[market],
                'counts': {key: 0 for key in DATASET_NAMES},
            })
            entry['counts'][kind] = len(records)
    return {
        'files': [{'kind': key, 'name': name, 'available': available[key]}
                  for key, name in FILES.items()],
        'scopes': sorted(scopes.values(), key=lambda s: (list(VEHICLES).index(s['model_code']), s['market'])),
        'notes': [
            'F70 feedback has country-specific comments; no German F70 verbatims are supplied.',
            'G60/G68 sales are one combined series total. Do not attribute the total to either model alone.',
            'The option catalog is a German-market F70 source. The curated feature labels need PM verification.',
            'Study mean ratings and US response fractions are different measures. 2030 sales are future-looking.',
        ],
    }


def load_scope(model_code: str, market_code: str) -> dict:
    if model_code not in VEHICLES or market_code not in MARKETS:
        raise ValueError('Choose a vehicle and market shown in the BMW data library.')
    data = _all()
    datasets = []
    for kind, mapping in data.items():
        records = mapping.get((model_code, market_code), [])
        if not records:
            continue
        sources = sorted({str(record['fields']['source_file']) for record in records})
        datasets.append({
            'kind': kind, 'type': DATASET_NAMES[kind],
            'name': f'BMW {model_code} {MARKETS[market_code]} · {DATASET_NAMES[kind]}',
            'sources': sources, 'records': records,
        })
    if not datasets:
        raise ValueError('No supplied BMW records match that vehicle and market.')
    return {'model_code': model_code, 'vehicle': VEHICLES[model_code],
            'market_code': market_code, 'market': MARKETS[market_code],
            'datasets': datasets, 'record_count': sum(len(d['records']) for d in datasets)}
